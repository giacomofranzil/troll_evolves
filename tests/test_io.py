"""Checks of the Excel layer, of the JSON contract and of the tracking import."""

from __future__ import annotations

import json

import pytest
from openpyxl import load_workbook

from hsmpace.core.analysis import analyse_sequence
from hsmpace.core.contract import case_from_dict, case_to_dict, report_to_dict
from hsmpace.core.model import harmonise_tandem_speeds
from hsmpace.core.studies import base_results, gap_vs_pacing, min_feasible_pacing, sequence
from hsmpace.core.tracking import parse_tracking
from hsmpace.example import example_case
from hsmpace.io_excel import ValidationError, read_case, write_case, write_results


def test_excel_round_trip_preserves_the_case(tmp_path):
    original = example_case()
    path = write_case(original, tmp_path / "case.xlsx")
    loaded = read_case(path)

    assert loaded.line.equipment == original.line.equipment
    assert loaded.line.sections == original.line.sections
    assert loaded.products == original.products
    assert loaded.settings == original.settings
    assert loaded.mill_type == original.mill_type == "hsm"
    assert loaded.utilities == original.utilities


def test_the_empty_template_has_the_sheets_and_headers(tmp_path):
    path = write_case(example_case(), tmp_path / "template.xlsx", include_data=False)
    wb = load_workbook(path)
    assert {"Guide", "Info", "Layout", "Sections", "Products", "PassSchedule", "Simulation", "Utilities"} <= set(
        wb.sheetnames
    )
    assert wb["Layout"]["A1"].value == "equipment_id"
    header = [c.value for c in wb["Sections"][1]]
    assert "d1_m" in header and "accel_mps2" in header
    assert wb["Layout"].max_row == 1, "the template must contain no data"
    guide = "\n".join(str(row[0].value or "") for row in wb["Guide"].iter_rows(min_col=1, max_col=1))
    assert "offline model TRoll" in guide
    assert "v* = sqrt" in guide
    assert "whichever coiler takes the strip" in guide
    assert "tandem slows down anyway" in guide
    assert "as late as possible" in guide
    assert "coiler_pattern" in guide
    assert "Several coilers" in guide
    assert "coilbox" in guide
    assert "bypass" in guide
    assert "occupy" in guide
    assert "relative change" in guide
    assert "Sheet Utilities" in guide
    assert "L/s" in guide
    assert wb["Utilities"]["A1"].value == "equipment_id"
    assert wb["Utilities"].max_row == 1, "the empty template must contain no recipes"


def test_json_round_trip_preserves_the_case():
    original = example_case()
    payload = json.loads(json.dumps(case_to_dict(original)))
    loaded = case_from_dict(payload)

    assert loaded.line.equipment == original.line.equipment
    assert loaded.products == original.products
    assert loaded.settings == original.settings
    assert loaded.mill_type == "hsm"
    assert payload["mill_type"] == "hsm"
    assert payload["contract_version"] == "1"
    assert payload["utilities"][0]["equipment_id"] == "DS1"


def test_parsing_errors_point_at_sheet_and_cell(tmp_path):
    path = write_case(example_case(), tmp_path / "broken.xlsx")
    wb = load_workbook(path)
    wb["Layout"]["C5"] = "twenty five"
    wb.save(path)

    with pytest.raises(ValidationError) as exc:
        read_case(path)
    issue = exc.value.issues[0]
    assert issue.sheet == "Layout"
    assert issue.cell == "C5"
    assert "is not a number" in issue.message


def test_model_errors_trace_back_to_the_pass_row(tmp_path):
    path = write_case(example_case(), tmp_path / "broken2.xlsx")
    wb = load_workbook(path)
    wb["PassSchedule"]["F4"] = 200.0  # h_out larger than h_in
    wb.save(path)

    with pytest.raises(ValidationError) as exc:
        read_case(path)
    reduction = [i for i in exc.value.issues if "invalid reduction" in i.message]
    assert reduction and reduction[0].sheet == "PassSchedule"
    assert reduction[0].cell == "B4"


def test_an_unsupported_schema_version_is_rejected(tmp_path):
    path = write_case(example_case(), tmp_path / "old.xlsx")
    wb = load_workbook(path)
    wb["Info"]["B2"] = "0"
    wb.save(path)

    with pytest.raises(ValidationError, match="schema_version"):
        read_case(path)


def test_a_section_event_beyond_the_length_is_rejected(tmp_path):
    path = write_case(example_case(), tmp_path / "section.xlsx")
    wb = load_workbook(path)
    ws = wb["Sections"]
    col = [c.value for c in ws[1]].index("d1_m") + 1
    ws.cell(row=4, column=col, value=999.0)
    wb.save(path)

    with pytest.raises(ValidationError, match="beyond the section length"):
        read_case(path)


def test_an_older_workbook_without_optional_columns_still_loads(tmp_path):
    """Columns added after the first schema version must not break existing files."""
    path = write_case(example_case(), tmp_path / "older.xlsx")
    wb = load_workbook(path)
    ws = wb["PassSchedule"]
    header = [c.value for c in ws[1]]
    ws.delete_cols(header.index("reversing_clearance_m") + 1)
    wb.save(path)

    loaded = read_case(path)
    assert all(p.reversing_clearance == 0.0 for p in loaded.products[0].passes)


def test_an_older_workbook_without_coiler_pattern_still_loads(tmp_path):
    """A single-coiler file written before the pattern key remains valid."""
    from dataclasses import replace

    from hsmpace.core.model import Line

    original = example_case()
    equipment = tuple(
        e for e in original.line.equipment if e.kind != "coiler" or e.id == "DC1"
    )
    one = replace(
        original,
        line=Line(equipment, original.line.sections),
        settings=replace(original.settings, coiler_pattern=()),
    )
    path = write_case(one, tmp_path / "no_pattern.xlsx")
    wb = load_workbook(path)
    ws = wb["Simulation"]
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == "coiler_pattern":
            ws.delete_rows(row)
            break
    wb.save(path)

    loaded = read_case(path)
    assert loaded.settings.coiler_pattern == ()
    assert loaded.piece_coiler_ids == ("DC1",) * loaded.settings.n_pieces


def test_two_coilers_in_excel_require_the_pattern(tmp_path):
    path = write_case(example_case(), tmp_path / "no_cycle.xlsx")
    wb = load_workbook(path)
    ws = wb["Simulation"]
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == "coiler_pattern":
            ws.cell(row=row, column=2).value = ""
            break
    wb.save(path)

    with pytest.raises(ValidationError, match="coiler_pattern"):
        read_case(path)


def test_a_misplaced_reversal_value_is_only_a_warning(tmp_path):
    """The delay belongs to the pass the reversal comes after, not to any pass."""
    path = write_case(example_case(), tmp_path / "misplaced.xlsx")
    wb = load_workbook(path)
    ws = wb["PassSchedule"]
    col = [c.value for c in ws[1]].index("reversing_delay_s") + 1
    ws.cell(row=14, column=col, value=9.0)  # last pass, no reversal follows it
    wb.save(path)

    loaded = read_case(path)
    assert any("no reversal follows" in w for w in loaded.warnings)


def test_writing_the_results(tmp_path):
    case, deviations = harmonise_tandem_speeds(example_case())
    base = base_results(case)
    results = sequence(case, base, case.settings.pacing)
    analyses = analyse_sequence(results, case.settings.gap_min, case.line)
    curve = gap_vs_pacing(case, base)
    best = min_feasible_pacing(case, base, curve)

    path = write_results(tmp_path / "out.xlsx", case, results, analyses, deviations, curve, best)
    wb = load_workbook(path)
    assert {"Summary", "Events", "Occupancy", "Utilities", "Gap", "PacingCurve", "MassBalance"} <= set(
        wb.sheetnames
    )
    assert wb["Events"].max_row > 10
    totals = [
        [c.value for c in row]
        for row in wb["Utilities"].iter_rows(min_row=2)
        if row[0].value == "TOTAL"
    ]
    assert any(r[2] == "water" and r[7] == "m3" and r[8] > 0 for r in totals)
    assert any(r[2] == "power" and r[7] == "kWh" and r[8] > 0 for r in totals)


def test_the_json_report_carries_segments_and_outcomes():
    case, deviations = harmonise_tandem_speeds(example_case())
    base = base_results(case)
    results = sequence(case, base, case.settings.pacing)
    analyses = analyse_sequence(results, case.settings.gap_min, case.line)

    report = report_to_dict(case, results, analyses, deviations)
    json.dumps(report)  # must be serialisable

    assert report["pieces"][0]["head_segments"][0]["t0_s"] == 0.0
    assert report["pieces"][0]["coiler_id"] in {"DC1", "DC2", "DC3"}
    assert report["pieces"][0]["length_kinematic_m"] == pytest.approx(
        report["pieces"][0]["length_geometric_m"], rel=1e-6
    )
    assert report["gaps"][0]["ok"] is True
    assert report["utilities"]["water_m3"] > 0.0
    assert report["utilities"]["power_kwh"] > 0.0


def test_an_older_workbook_without_occupy_columns_still_loads(tmp_path):
    path = write_case(example_case(), tmp_path / "no_occupy.xlsx")
    wb = load_workbook(path)
    ws = wb["Layout"]
    header = [c.value for c in ws[1]]
    for name in ("occupy_after_m", "occupy_before_m", "occupy"):
        ws.delete_cols(header.index(name) + 1)
        header = [c.value for c in ws[1]]
    wb.save(path)

    loaded = read_case(path)
    ds1 = loaded.line.get("DS1")
    assert ds1.occupy is None
    assert ds1.occupy_before == 0.0
    assert not ds1.occupies


def test_an_older_workbook_without_the_product_coilbox_tick_still_loads(tmp_path):
    path = write_case(example_case(), tmp_path / "no_cb_tick.xlsx")
    wb = load_workbook(path)
    ws = wb["Products"]
    header = [c.value for c in ws[1]]
    ws.delete_cols(header.index("coilbox") + 1)
    wb.save(path)

    loaded = read_case(path)
    assert all(p.use_coilbox is None for p in loaded.products)


def test_coilbox_bypass_round_trips_through_excel(tmp_path):
    from dataclasses import replace

    original = example_case()
    product = replace(original.products[0], use_coilbox=False)
    original = replace(original, products=(product,) + original.products[1:])
    path = write_case(original, tmp_path / "bypass.xlsx")
    loaded = read_case(path)
    assert loaded.products[0].use_coilbox is False


def test_tracking_csv_import():
    rows = [
        "piece_id,time_s,head_m,tail_m",
        "A1,0.0,0.0,-10.5",
        "A1,1.0,1.2,-9.3",
        "A2,0.0,0.0,-10.5",
    ]
    series = parse_tracking(rows)
    assert len(series) == 2
    first = next(s for s in series if s.piece_id == "A1")
    assert first.t == [0.0, 1.0]
    assert first.head == [0.0, 1.2]
    assert first.shift(10.0).t == [10.0, 11.0]


def test_tracking_without_the_tail_column():
    series = parse_tracking(["piece_id,time_s,head_m", "A1,0.0,5.0"])
    assert series[0].tail == [5.0]


def test_tracking_with_missing_columns():
    with pytest.raises(ValueError, match="missing columns"):
        parse_tracking(["piece,time", "A1,0"])


def test_an_older_workbook_without_the_utilities_sheet_still_loads(tmp_path):
    path = write_case(example_case(), tmp_path / "no_utilities.xlsx")
    wb = load_workbook(path)
    del wb["Utilities"]
    wb.save(path)

    loaded = read_case(path)
    assert loaded.utilities == ()


def test_json_without_utilities_defaults_to_empty():
    payload = case_to_dict(example_case())
    del payload["utilities"]
    loaded = case_from_dict(payload)
    assert loaded.utilities == ()


def test_json_without_mill_type_defaults_to_hsm():
    payload = case_to_dict(example_case())
    del payload["mill_type"]
    payload.get("info", {}).pop("mill_type", None)
    loaded = case_from_dict(payload)
    assert loaded.mill_type == "hsm"


def test_unknown_mill_type_is_rejected():
    from dataclasses import replace

    from hsmpace.core.model import validate_case

    case = replace(example_case(), mill_type="plate")
    messages = [p.message for p in validate_case(case)]
    assert any("mill_type" in m and "plate" in m for m in messages)


def test_troll_xml_import_is_a_documented_boundary():
    from hsmpace.io.troll_xml import case_from_troll

    with pytest.raises(NotImplementedError, match="TRoll XML"):
        case_from_troll("dump.xml")


def test_io_excel_shim_still_imports():
    from hsmpace.io.excel import read_case as read_new
    from hsmpace.io_excel import read_case as read_shim

    assert read_new is read_shim
